#!/usr/bin/env python3
"""Build the mlx5 HWS analysis Word report from the simulation CSV and PNGs."""
import csv
import os
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
FIG = os.path.join(ROOT, "figures")
CSV = os.path.join(ROOT, "sim", "out", "summary.csv")
OUT = os.path.join(ROOT, "mlx5-hws-model.docx")
FONT = "微软雅黑"


def set_run(run, size=11, bold=False, color=None):
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_p(doc, text, size=11, bold=False, center=False, space=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run(run, size=size, bold=bold)
    return p


def add_h(doc, text, level):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run(run, size={1: 18, 2: 15, 3: 13}[level], bold=True, color=(27, 58, 75))
    return h


def add_fig(doc, name, caption):
    path = os.path.join(FIG, name)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(6.35))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(12)
    run = cap.add_run(caption)
    set_run(run, size=9, color=(60, 70, 80))


def shade_header(cell):
    tc = cell._tePr if hasattr(cell, "_tePr") else cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = tcPr.makeelement(qn("w:shd"), {
        qn("w:fill"): "1B3A4B",
        qn("w:val"): "clear",
    })
    tcPr.append(shd)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        set_run(run, size=9, bold=True, color=(255, 255, 255))
        shade_header(cell)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.rows[r + 1].cells[c]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(val))
            set_run(run, size=8)
    doc.add_paragraph()


def load():
    with open(CSV, newline="") as f:
        rows = list(csv.DictReader(f))
    return {(r["exp"], r["data"]): r for r in rows}


def pct(x):
    return f"{float(x) * 100:.4f}%"


def rel(emp, theory):
    theory = float(theory)
    emp = float(emp)
    if theory == 0:
        return "公式为 0"
    return f"{(emp - theory) / theory * 100:+.2f}%"


def main():
    by = load()
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    add_p(doc, "mlx5 硬件流表（HWS）结构、流水线与插入上界", size=22, bold=True, center=True, space=4)
    add_p(doc, "从本仓库 drivers/net/ethernet/mellanox/mlx5/core/steering/hws 反推", size=12, center=True, space=2)
    add_p(doc, "2026-10-28  Asia/Shanghai", size=11, center=True, space=12)
    add_p(doc, "本文对照的是驱动里的表几何和一份按该几何写成的用户态模型。环境里没有 ConnectX，文中的秒数是模型程序自己的运行时间，不是网卡插入速率。")

    add_h(doc, "1  结论", 1)
    add_p(doc, "以太网方向上有三个并列的域：NIC_RX、NIC_TX、FDB。fs_core 为它们各建一个根 namespace。NIC_RX 和 NIC_TX 的命令留在固件流表。FDB 根可以切到 smfs 或 hmfs。后面关于 RTC、STE 和插入上界的部分，写的是切到 hmfs 之后的 FDB 域。FDB 内部的 RX、TX 两套 RTC 是这一个域的两个方向。")
    add_p(doc, "HWS 把一张 FDB 流表拆成按优先级串起来的 matcher。每个 matcher 有一套匹配模板（key）和若干动作模板（action）。同一条规则建两套 RTC：RX 服务线侧进入的包，TX 服务 vport 发出的包。大表在主 RTC 后面再挂一张碰撞 RTC，优先级与主 matcher 相同。")
    add_p(doc, "存储单位是 STE。RESOURCE_MODE_RULE 下，调用者给出的 num_log 是哈希行数的对数。num_log = 24 时主表是 2^24 行、每行 4 个 STE，一共 2^26 个槽。碰撞表是 2^19 行、每行 16 个槽。2^24 条 key 均匀放入时，主表大约四分之三的槽是空的。这些空槽用来吸收同一行上的哈希起伏，使溢出比例停在千分之几。")
    add_p(doc, "插入由专用发送队列上的 GTA WQE 完成，opcode 为 0x2c。硬件只在哈希选中的那一行里找空 STE。行满时软件再向碰撞 RTC 发一条 WQE。成功路径的 WQE 数只有 1 或 2，与表中已有规则数无关。千万条每秒这个量级要靠设备把这种定长 WQE 在多条队列上打满；规则变多不会把单次插入变成随表长增长的搜索。")
    a = by[("A", "uniform")]
    b = by[("B", "uniform")]
    d = by[("Dtight", "uniform")]
    add_p(doc, f"均匀哈希、1,000 万条、驱动几何：溢出 {pct(a['over_frac'])}，公式 {pct(a['theory_over'])}，相对误差 {rel(a['over_frac'], a['theory_over'])}，失败 0。配满 2^24 条时溢出 {pct(b['over_frac'])}，公式 {pct(b['theory_over'])}，相对误差 {rel(b['over_frac'], b['theory_over'])}，失败仍为 0。把主表收成 2^22 行后再放入同样的 2^24 条 key，失败率 {pct(d['fail_frac'])}，公式 {pct(d['theory_fail'])}。模型与驱动几何在均匀负载上对得上；把行数按“槽位总数等于 key 数”收紧之后，失败率会离开设计点。")

    add_h(doc, "2  源码地图", 1)
    add_p(doc, "分析只读本仓库，不改驱动。HWS 入口在 steering/hws。与 fs_core 的衔接在 fs_hws.c：只有 table_type 为 FS_FT_FDB 的 namespace 会走 mlx5hws_table_create，其它 namespace 仍走固件流表命令。")
    add_table(doc,
              ["对象", "文件", "在硬件侧对应什么"],
              [
                  ["context", "context.c", "每设备一个上下文：STC 池、definer 缓存、发送队列"],
                  ["table", "table.c", "一张 FDB flow table，带 level 和默认 miss"],
                  ["matcher", "matcher.c", "一张 RTC 哈希表，按 priority 串进表"],
                  ["definer", "definer.c", "从 match criteria 选出要写入 STE tag 的字段"],
                  ["action / STC", "action.c pat_arg.c", "改头、计数、vport、跳表等动作"],
                  ["rule", "rule.c", "一条规则：tag、动作、rtc_0 / rtc_1"],
                  ["send", "send.c", "GTA WQE、门铃、CQ 完成；满行则重试碰撞 RTC"],
                  ["bwc", "bwc.c", "按规则数扩行，70% 水位触发 rehash"],
                  ["fs 衔接", "fs_hws.c", "flow group 变成 bwc matcher，目的地变成 action"],
              ])

    add_h(doc, "3  域、namespace、表、key 与 action", 1)
    add_h(doc, "3.1  三个域", 2)
    add_p(doc, "NIC_RX、NIC_TX、FDB 是三套并列的以太网流表域。驱动里有两层同名枚举。固件侧是 fs_flow_table_type：FS_FT_NIC_RX = 0，FS_FT_NIC_TX = 1，FS_FT_FDB = 4。软件转向侧是 steering/sws/mlx5dr.h 里的 mlx5dr_domain_type，三个取值也是 NIC_RX、NIC_TX、FDB。fs_core 初始化时各建一个根：init_root_ns() 得到 steering->root_ns，表类型 NIC_RX；init_egress_root_ns() 得到 egress_root_ns，表类型 NIC_TX；设备是 eswitch manager 且 FDB 能力打开时，init_fdb_root_ns() 得到 fdb_root_ns。SET_FLOW_TABLE_ROOT 按 table_type 分别挂根，NIC_TX 的入口在 mlx5_fs_cmd_set_tx_flow_table_root()。")
    add_fig(doc, "21_domains.png", "图 21  NIC_RX、NIC_TX、FDB 三个域。前两个在本树走固件流表，FDB 可以切到 HWS")
    add_table(doc,
              ["域", "根", "本树的命令", "方向", "namespace", "典型终结"],
              [
                  ["NIC_RX", "root_ns", "固件 DMFS", "只建 RX 半边",
                   "BYPASS、RX MACSEC、LAG、OFFLOADS、ETHTOOL、KERNEL、LEFTOVERS、ANCHOR",
                   "TIR，再由 TIR 选 RQ"],
                  ["NIC_TX", "egress_root_ns", "固件 DMFS", "只建 TX 半边",
                   "EGRESS、EGRESS_IPSEC、EGRESS_MACSEC",
                   "改头、L2 到隧道的封装，然后离开 function"],
                  ["FDB", "fdb_root_ns", "dmfs、smfs 或 hmfs", "RX 与 TX 各一套 STE",
                   "FDB 根、FDB_BYPASS，以及运行时的 TC chain",
                   "vport、丢弃、跳表、封装"],
              ])
    add_p(doc, "表 4  三个以太网域。RDMA、sniffer、port select、vport ACL 另外还有自己的根和表类型，不在这三列里。", size=9)
    add_p(doc, "mlx5dr_domain_create() 按类型填半边。NIC_RX 要求接收方向的软件转向能力，默认 ICM 用 nic_rx_drop_address，只初始化 info.rx。NIC_TX 用 nic_tx_allow_address，只初始化 info.tx。FDB 要求 eswitch manager，RX、TX 的默认 ICM 取管理 vport 的两个地址，丢弃地址用 esw_rx_drop 和 esw_tx_drop。表和规则跟着分叉：NIC 域的一张表只有一个方向的锚点，FDB 的一张表在 dr_table_init_fdb() 里把 rx 和 tx 都建出来。FDB 规则还可以按 flow_source 和 source_port 跳过其中一边，这段在 dr_rule_skip()。")
    add_p(doc, "上一版把域收成 FDB，是因为跟上了 HWS 的创建入口。mlx5_flow_namespace_set_mode() 在根的 table_type 不是 FS_FT_FDB 时返回 -EOPNOTSUPP。全库只有 esw_create_offloads_fdb_tables() 调用它，传入的 namespace 是 MLX5_FLOW_NAMESPACE_FDB。hmfs 因此只在 FDB 根上打开 mlx5hws_context。mlx5_cmd_hws_create_flow_table() 看到其它 table_type 直接返回 -EOPNOTSUPP。enum mlx5hws_table_type 里只有 MLX5HWS_TABLE_TYPE_FDB。smfs 的 mlx5_cmd_dr_create_ns() 调用 mlx5dr_domain_create() 时写死 MLX5DR_DOMAIN_TYPE_FDB。NIC_RX 和 NIC_TX 的创建分支在 dr_domain.c 里是完整的，fs 层没有走到。这两个根的 cmds 一直是 mlx5_fs_cmd_get_default() 返回的固件命令。")
    add_p(doc, "HWS 源码还留着 NIC 域的口子。表类型不是 FDB 且带了 optimize_flow_src 时，matcher.c 打印 “NIC domain doesn't support flow_src”。table.c 在非 FDB 时用 nic_ft.max_level 填下一张表的 level。mlx5hws_table_get_fw_ft_type() 对非 FDB 返回 -EOPNOTSUPP。本树的 HWS 不会打开 NIC_RX 或 NIC_TX。")
    add_p(doc, "FDB 域内部还有两个固件资源类型：FS_FT_FDB_RX 与 FS_FT_FDB_TX。hws_matcher_create_rtc() 在表类型为 FDB 时创建第二套 RTC，STE 区间记在 ste_0_base 与 ste_1_base。这两套 RTC 是 FDB 这一个域的两个方向，和 NIC_RX、NIC_TX 两个域并列。context 初始化还要求 eswitch manager、WQE 更新、8DW STE、按哈希和按偏移两种 RTC 更新模式，以及 SELECT definer。不满足这些能力时，上下文不会置上 HWS_SUPPORT。")
    add_h(doc, "3.2  namespace 里有几张表", 2)
    add_p(doc, "namespace 挂在某一个域的根下面。mlx5_get_flow_namespace() 把 BYPASS 到 ANCHOR 这 8 个类型交给 root_ns，把 EGRESS 三个类型交给 egress_root_ns，把 FDB 和 FDB_BYPASS 交给 fdb_root_ns。HWS 创建路径只接受 FS_FT_FDB，所以 hmfs 的 matcher 只出现在 FDB 域里。NIC_RX 和 NIC_TX 的表由固件流表命令创建。")
    add_p(doc, "FDB namespace 内部的表数随 offload 变化，不是一个写死的整数。fs.h 里的优先级梯子是：FDB_BYPASS_PATH、FDB_CRYPTO_INGRESS、FDB_TC_OFFLOAD、FDB_FT_OFFLOAD、FDB_TC_MISS、FDB_BR_OFFLOAD、FDB_SLOW_PATH、FDB_CRYPTO_EGRESS、FDB_PER_VPORT。esw_create_offloads_fdb_tables() 至少创建两张表：慢路径表（优先级 FDB_SLOW_PATH）和 TC miss 表（优先级 FDB_TC_MISS）。TC chain 在运行时继续按 level 加表。每张 mlx5_flow_table 对应一张 mlx5hws_table。")
    add_p(doc, "一张 HWS 表里的 matcher 数等于这张表上的 flow group 数，再加每个大 matcher 的一张碰撞表。碰撞 matcher 使用与主 matcher 相同的 priority，创建时插到主 matcher 后面，查找顺序是先主 RTC、再碰撞 RTC。")
    add_h(doc, "3.3  每张表的 key 和 action", 2)
    add_p(doc, "key 不是全 FDB 共用的一张五元组。fs_hws 把每个 flow group 的 match_criteria 交给 bwc matcher，definer 再把置位字段收成 STE tag。普通 tag 32 字节，jumbo tag 44 字节。不同 group 的 key 不同。慢路径表上驱动写死的几张 group 是：")
    add_table(doc,
              ["flow group", "key", "action"],
              [
                  ["send-to-vport", "source_sqn，必要时再加源 eswitch 的 vhca id", "转到对应 vport"],
                  ["meta send-to-vport", "metadata_reg_c_0 的 vport 元数据，以及 reg_c_1 的隧道位", "转到 vport"],
                  ["miss", "外层 DMAC 的组播位（dmac[0] = 0x01）", "单播和组播 miss 都转到管理 vport"],
              ])
    add_p(doc, "业务规则（TC、桥、封装）的 key 由当时的 flow spec 决定，常见字段来自 definer 能选的外层/内层地址、端口、VLAN、隧道和寄存器。action 来自动作模板，类型包括丢弃、跳表、计数、tag、改头、vport、压弹 VLAN、封装与解封装、插入或删除头、meter、采样和目的地数组。终结到 vport 的动作在 fs_hws.c 里通过 mlx5hws_action_create_dest_vport 创建。")
    add_p(doc, "上面这张慢路径表是 switchdev 的 FDB。NIC_RX 的 KERNEL namespace 由 en_fs.c 建 TTC，两种模式的本机网口都会建。fs_ttc.c 用外层 ethertype、IP 版本和 L4 协议做 key，目的地类型是 TIR，tir_num 指向 RSS 或 direct TIR。HWS 动作枚举里没有 TIR。")

    add_h(doc, "3.4  legacy 和 switchdev", 2)
    add_p(doc, "两个模式都保留 NIC_RX、NIC_TX、FDB 三个根。eswitch 创建时 mode 写成 MLX5_ESWITCH_LEGACY，devlink 切到 switchdev 之后变成 MLX5_ESWITCH_OFFLOADS。变的是每个根里实际创建的流表和规则，以及 FDB 能不能离开固件流表。ACL 另有表类型 FS_FT_ESW_INGRESS_ACL 和 FS_FT_ESW_EGRESS_ACL，按 vport 各有一个根，两种模式都配，规则不一样。没开 eswitch 的纯网卡不走这两条路径，只有 NIC_RX 和 NIC_TX。")
    add_h(doc, "3.5  legacy", 2)
    add_fig(doc, "22_legacy.png", "图 22  legacy。FDB 按目的 MAC 选 vport，收包仍走该 function 自己的 NIC_RX")
    add_p(doc, "esw_legacy_enable() 调用 esw_create_legacy_fdb_table()。表建在 mlx5_get_fdb_sub_ns(dev, 0)，也就是 FDB_TC_OFFLOAD 优先级下的第 0 个 chain namespace，优先级 LEGACY_FDB_PRIO = 1。命令留在固件流表，set_mode() 不会被调用，所以没有 smfs 域，也没有 hmfs context。")
    add_table(doc,
              ["表", "位置", "key", "action"],
              [
                  ["主 FDB，地址组", "chain 0，prio 1，索引 0 到 size-3", "外层 DMAC 全匹配", "转到拥有该 MAC 的 vport"],
                  ["主 FDB，allmulti", "同一张表，索引 size-2", "外层 DMAC 的组播位", "转到打开 allmulti 的 vport"],
                  ["主 FDB，混杂", "同一张表，索引 size-1", "source_port = uplink", "转到打开混杂的 vport"],
                  ["VEPA", "同一 namespace，prio 0，只在打开 VEPA 时创建", "source_port = uplink", "转到主 FDB；另一条把其余流量转到 uplink vport"],
                  ["ingress ACL", "非管理 vport 的 ESW_INGRESS", "CVLAN 标签、SMAC", "允许合法源，其余丢弃"],
                  ["egress ACL", "非管理 vport 的 ESW_EGRESS", "外层 CVLAN VID", "允许指定 VLAN，其余丢弃并计数"],
                  ["NIC_RX", "PF，以及客户机里的每个 VF", "KERNEL 的 L2、TTC", "本 function 的 TIR"],
                  ["NIC_TX", "同上", "EGRESS、IPsec、MACsec", "改头或封装后离开该 function"],
              ])
    add_p(doc, "表 5  legacy 实际创建的表。管理 vport 不建 legacy ACL。", size=9)
    add_p(doc, "线侧报文先在这张 FDB 里按目的 MAC 选 vport，然后进入那个 function 自己的 NIC_RX，再由 TIR 选 RQ。VF 发出的报文先过该 VF 的 NIC_TX，再回到同一张 FDB。地址规则在 __esw_fdb_set_vport_rule() 里写，目的地类型是 VPORT。")
    add_h(doc, "3.6  switchdev", 2)
    add_fig(doc, "23_switchdev.png", "图 23  switchdev。FDB 可以切到 HWS，代表口从 NIC_RX 的 OFFLOADS 表收包")
    add_p(doc, "esw_offloads_enable() 进入 esw_offloads_steering_init()。FDB 表由 esw_create_offloads_fdb_tables() 创建，这里才会调用 mlx5_flow_namespace_set_mode()，所以 smfs 或 hmfs 只出现在这个模式。慢路径表和 TC miss 表挂在 FDB 根上，优先级分别是 FDB_SLOW_PATH 和 FDB_TC_MISS。TC chain 和 FT chain 的 namespace 在 fs_core 初始化时已经排好，表在运行时按 chain 加进去。")
    add_table(doc,
              ["表", "位置", "key", "action"],
              [
                  ["慢路径 send-to-vport", "FDB 根，FDB_SLOW_PATH", "source_sqn，合并 eswitch 时再加 vhca id", "转到 vport"],
                  ["慢路径 meta send-to-vport", "同一张表", "reg_c_0 的 vport 元数据，reg_c_1 的隧道位", "转到 vport"],
                  ["慢路径 miss", "同一张表", "外层 DMAC 的组播位", "转到管理 vport"],
                  ["TC miss", "FDB 根，FDB_TC_MISS", "空表，用来把后面的优先级接上", "默认 miss 到下一张表"],
                  ["TC / FT chain", "FDB_TC_OFFLOAD、FDB_FT_OFFLOAD 的 chain", "当时的 flow spec", "封装、改头、跳表、vport、丢弃"],
                  ["ft_offloads", "NIC_RX 的 OFFLOADS，prio 1", "source_port，或 reg_c_0", "转到代表口的 root 或 TTC"],
                  ["ft_offloads 丢弃", "同一张表的最后一条", "空匹配", "丢弃"],
                  ["restore", "NIC_RX 的 OFFLOADS，需要 reg_c1 环回能力", "reg_c_0", "把 reg_c_1 拷到 reg_b"],
                  ["代表口 TTC", "NIC_RX 的 KERNEL", "外层 ethertype、L4", "该代表口的 TIR"],
                  ["上行代表口 root", "NIC_RX 的 OFFLOADS，level 1", "空表", "miss 接到 ethtool 或 TTC"],
                  ["ingress ACL", "uplink、管理 vport，以及后来启用的 vport", "未打标签的包，或全匹配", "把 vport 元数据写入 reg_c_0；可选压上 prio tag"],
                  ["egress ACL", "同上", "全匹配，或 prio tag", "转到 vport，或弹出 prio tag"],
              ])
    add_p(doc, "表 6  switchdev 实际创建的表。ft_offloads 和 restore 只在 esw_offloads_steering_init() 里出现。", size=9)
    add_p(doc, "线侧报文先过 FDB。命中后目的地是 vport。代表口加载时，en_rep.c 在 ft_offloads 上为这个 vport 加一条规则，目的地是一张流表：上行口是 OFFLOADS 里的空 root，其它代表口直接是自己的 TTC。TTC 再把报文送到该代表口的 TIR。发送方向先过代表口的 NIC_TX，再进入 FDB 的 TX 侧。")
    add_p(doc, "两边对照下来：legacy 的 FDB 是一张 L2 MAC 表，switchdev 的 FDB 是慢路径加上 TC chain，并且只有后者能切到 hmfs。legacy 的 NIC_RX 属于 PF 和客户机里的 VF；switchdev 在 NIC_RX 的 OFFLOADS namespace 里多出 ft_offloads 和 restore，收包的是 hypervisor 上的代表口。ACL 的根都在，legacy 用它做 spoof check 和 VLAN 过滤，switchdev 用它给报文盖上 vport 元数据。")

    add_h(doc, "4  存储、插入、查找", 1)
    add_fig(doc, "01_objects.png", "图 1  context、table、matcher、RTC、STE、STC 的关系")
    add_h(doc, "4.1  怎么存", 2)
    add_p(doc, "RTC 的 log_hash_size 来自 sz_row_log，log_depth 来自 sz_col_log。RESOURCE_MODE_RULE 且按哈希插入时，这两个字段共用 union 的第一个字节：调用者写入 rule.num_log，驱动只把第二个字节写成深度。大表（num_log > 10）的深度固定为 MLX5HWS_MATCHER_ASSURED_MAIN_TBL_DEPTH，也就是 2，每行 4 个 STE。STE 池的对数大小是行对数加深对数。")
    add_p(doc, "碰撞表在 hws_matcher_create_col_matcher() 里生成。行对数等于 num_log 减 MLX5HWS_MATCHER_ASSURED_ROW_RATIO（5），深度对数为 4，每行 16 个 STE。注释里的 3% 来自 log2(0.03) ≈ 5：碰撞行数约为主表行数的 1/32。")
    add_fig(doc, "07_geometry.png", "图 2  左边是驱动几何：2^24 行、每行 4 个槽。右边是紧凑几何：2^22 行、每行 4 个槽，槽位总数等于 2^24 条 key")
    add_p(doc, "后文的“驱动几何”就是左图这套尺寸。num_log = 24 时，驱动把主表配成 2^24 行、每行 4 个 STE，碰撞表配成 2^19 行、每行 16 个 STE。它不是“槽位个数等于 key 个数”。右图那套称为紧凑几何，只在对照实验里使用。")
    add_h(doc, "4.2  怎么插入，时间上界从哪来", 2)
    add_fig(doc, "05_insert.png", "图 3  插入路径。满行只再发一条碰撞 WQE")
    add_p(doc, "rule.c 把主 RTC 填进 rtc_0 / rtc_1，把碰撞 RTC 填进 retry_rtc_0 / retry_rtc_1。flow_source 为 WIRE 或 VPORT 时可以跳过另一方向。send.c 在 CQE 报失败且 retry_id 非 0 时，把同一条规则再投递到碰撞 RTC。碰撞 RTC 再失败，规则状态变为失败，不再发第三条 WQE。")
    add_p(doc, "因此一次成功插入的设备操作是 1 次或 2 次定长 WQE（控制段 48 字节，数据段 64 字节）。比较次数被行深卡住：主行 4，碰撞行 16。这个上界不随 n 增长。BWC 在规则数达到 2^size_log 的 70% 时把行数加一档，扩的是行，不是把深度吃满。热路径按队列加锁，队列之间可以并行提交。")
    add_h(doc, "4.3  怎么查找", 2)
    add_fig(doc, "06_lookup.png", "图 4  报文查找。碰撞表在同一优先级上接在主表后面")
    add_p(doc, "hws_matcher_connect() 按 priority 把 matcher 串起来。前一张的 end FT 指向下一张的 RTC；最后一张指向 miss 表。碰撞 matcher 与主 matcher 优先级相等，插在后面，所以主 RTC 未命中会进入碰撞 RTC。命中后执行 STE 里指向的 STC。整条 matcher 链都未命中时，FDB 默认 miss 把报文转到 eswitch manager vport。")

    add_h(doc, "5  流水线，以及和 TIR、TIS、RQ、SQ、CQ、EQ 的层次", 1)
    add_fig(doc, "02_rx_pipeline.png", "图 5  RX。FDB 域在前，目的 function 的 NIC_RX 域在后，TIR 和 RQ 属于 NIC_RX")
    add_fig(doc, "03_tx_pipeline.png", "图 6  TX。SQ 和 TIS 之后先过 NIC_TX 域，再进 FDB 域的 TX RTC")
    add_fig(doc, "04_fdb_dual_rtc.png", "图 7  一条 FDB 规则写入 RX、TX 两套 STE，各自带碰撞 RTC")
    add_p(doc, "图 5 和图 6 画的是 switchdev。RX 的层次是：上行口，FDB 域（hmfs 时走 RTC RX 和 STC），目的 vport，NIC_RX 的 OFFLOADS 表，代表口的 TTC，TIR，RQ，CQ，EQ。legacy 的路径在图 22：FDB 按 DMAC 选 vport，然后进入该 function 自己的 NIC_RX。没有 eswitch 时，线侧报文直接进入 NIC_RX 域。")
    add_p(doc, "TIS 是发送侧的传输接口。主机把数据包 WQE 放进 SQ 并敲门铃之后，报文先进入该 function 的 NIC_TX 域。switchdev 下下一站是 FDB 域的 TX 侧，hmfs 时用 TX RTC 做哈希。legacy 下下一站是同一张按 DMAC 转发的 FDB。没有 eswitch 时，NIC_TX 之后直接出端口。")
    add_p(doc, "规则下发不走数据包 RQ/SQ。context 按队列数打开一组 HWS SQ/CQ，上面只跑 GTA WQE。完成由驱动轮询这条 CQ 取回；EQ 处于通知层，插入快路径不依赖 EQ 中断。所以图里有两条队列：一条搬包，一条搬 STE。")

    add_h(doc, "6  数学模型", 1)
    add_p(doc, "记主表行数 R = 2^r，每行容量 d = 4。碰撞表行数 Rc = R / 32，每行容量 dc = 16。插入 n 条互异规则，行哈希均匀时，每行到达数 X 服从 Poisson(λ)，λ = n / R。R 在实验里是 2^24 或 2^22，泊松近似足够。")
    add_p(doc, "主表放不下的比例 f(n) = E[max(X − d, 0)] / λ。这些规则再经独立哈希进入碰撞表，λc = n · f(n) / Rc，Y ~ Poisson(λc)。失败比例 p_fail(n) = f(n) · E[max(Y − dc, 0)] / λc。λc 为 0 时失败比例为 0。")
    add_p(doc, "主表槽位利用率 u_main = 主表实际占用槽数 / (4R)。设计点 n = R 时，若溢出可以忽略，u_main 约为 25%。成功插入的 GTA 步数 S 属于 {1, 2}：主行未满为 1，否则为 2。查找比较次数不超过 d，落到碰撞行时不超过 d + dc。")
    add_p(doc, "单队列规则速率可以写成 R_wqe / (1 + f(n))，其中 R_wqe 是该队列每秒能完成的 GTA WQE 数，Q 条队列近似乘 Q。f(n) 在驱动几何的配满点约为 0.44%，第二步带来的额外 WQE 不到百分之一。本文没有测量 R_wqe。上界成立时，表从空涨到 2^24 条不会把单次插入从常数次 WQE 变成线性扫描。")

    add_h(doc, "7  千万级验证", 1)
    add_p(doc, "程序是 mlx5-hws-analysis/sim/hws_sim.c。主表和碰撞表都按槽位数组逐条写入 64 位指纹。空槽为 0。插入之后按同一哈希把每条 key 再查一遍：成功的必须命中，失败的必须查不到，比较次数不得超过 4 + 16。全部实验的 lookup_bad 都是 0。这些秒数是模型程序的运行时间，不是网卡插入速率。")
    add_p(doc, "驱动几何：主表 2^24 行、每行 4 个 STE，碰撞表 2^19 行、每行 16 个 STE。配满是把 16,777,216 条 key 放进这 2^24 行，平均每行 1 条，λ = 1。紧凑几何：主表 2^22 行、每行 4 个 STE，碰撞表 2^17 行、每行 16 个 STE。同一批 16,777,216 条 key 放进 2^22 行时，平均每行 4 条，λ = 4。图上的 25%、50%、70%、90%、100% 都是插入条数除以 2^24，不是槽位被填满的比例。100% 在驱动几何上是平均每行 1 条；同一批 key 放到紧凑几何上仍标成 100%，但平均每行变成 4 条。70% 是 bwc.c 开始扩行的水位。")
    add_table(doc,
              ["图上的刻度", "插入条数", "主表", "这一档在看什么"],
              [
                  ["1000万均匀", "10,000,000", "2^24 行 × 4", "驱动几何，λ = 0.596"],
                  ["1677万，2^24 行", "16,777,216", "2^24 行 × 4", "配满，λ = 1"],
                  ["2000万", "20,000,000", "2^24 行 × 4", "超过配满，λ = 1.192"],
                  ["另一哈希", "10,000,000", "2^24 行 × 4", "换一套混合常数"],
                  ["重尾全字段", "10,000,000", "2^24 行 × 4", "Zipf 源地址再和序号一起混合"],
                  ["CTU 采样", "10,000,000", "2^24 行 × 4", "按真实五元组的经验频率采样"],
                  ["CTU 全量", "1,602,033", "2^24 行 × 4", "原始互异五元组"],
                  ["无副表", "10,000,000", "2^24 行 × 4", "去掉碰撞表"],
                  ["粗哈希", "10,000,000", "2^24 行 × 4", "行号只由 Zipf 组决定"],
                  ["1677万，2^22 行", "16,777,216", "2^22 行 × 4", "紧凑几何，λ = 4"],
                  ["25% 到 100%", "该比例 × 2^24", "看该图标题", "负载 = 条数 / 2^24"],
              ])
    add_p(doc, "表 0  实验图横轴刻度。1677 万是 2^24 = 16,777,216 的四舍五入写法，程序里用的是精确整数。", size=9)
    add_h(doc, "7.1  数据集", 2)
    add_p(doc, "均匀集：每条规则一个互异序号，经 64 位混合后取低位作为行号。另用第二套混合常数复跑 1,000 万条，用来看结论是否粘在某一个混合函数上。")
    add_p(doc, "重尾有两套。zipf_mixed：源组服从 Zipf，指数 1.2，组数 2^20，再与唯一序号一起混合后取行。这代表五元组字段都进了哈希，重尾被打散。zipf_coarse：行号只由 Zipf 源组决定，碰撞行仍用唯一序号的独立哈希。这代表 matcher 的 key 很短，或者哈希没有把热点字段打散。")
    add_p(doc, "真实数据用 CTU-13 的 capture20110810.binetflow（捷克理工大学 MCFP 公开的 botnet 流记录，2015-01-30 发布的文件）。原始 2,824,636 行，去掉非 IPv4 之后得到 1,602,033 个互异五元组指纹。全量作为 TraceFull。不足 1,000 万条，因此再按该文件里源地址、目的地址、协议、源端口、目的端口的经验频率独立采样，生成 1,000 万个互异组合，记为 ctu_resample。采样组合是经验分布的扩展，逐条列出的部分才是原始流。")
    add_h(doc, "7.2  结果", 2)
    add_p(doc, "图 8 的横轴是一次插入实验。左图四档从左到右：1000 万条均匀、1677 万条配满、2000 万条、CTU 采样 1000 万条，主表都是 2^24 行。纵轴是主表 4 个槽放不下的 key 占该档条数的百分比。右图在这四档后面加了第五档：同样 1677 万条放进 2^22 行。右图纵轴是两级都放不下的 key，也就是插入失败率。左图不画第五档，是因为 19.53% 会把前四档千分之几的柱子压成一条线。")
    add_fig(doc, "08_overflow_fail.png", "图 8  左：主表溢出率，横轴四档都是 2^24 行。右：插入失败率，最后一档是 2^22 行")
    add_p(doc, "前四档仿真柱和 Poisson 公式柱贴在一起，失败率是 0。第五档失败率升到约 7.07%，公式给出同一高度。驱动几何把溢出留在主表的千分之几，紧凑几何把同一批 key 的失败率抬到百分之七。")
    add_p(doc, "小结：均匀和真实五元组在驱动几何上失败为 0；把行数收成 2^22，失败率离开设计点。")

    add_p(doc, "图 9 的横轴十档从左到右是：1000 万均匀、1677 万配满、2000 万、1000 万另一套哈希、1000 万重尾但全字段进哈希、CTU 采样 1000 万、CTU 全量 160 万、1000 万去掉碰撞表、1677 万放进 2^22 行、1000 万粗哈希。十档里只有“2^22 行”的主表是 2^22 行 × 4 槽，其余九档都是 2^24 行 × 4 槽。“无副表”去掉了碰撞表，其余九档都带碰撞表。纵轴三段叠在一起，加起来是 100%：绿色是主行有空槽、只发 1 次 WQE，橙色是主行已满、再发 1 次碰撞 WQE，红色是碰撞行也满、插入失败。图例在图的右侧。")
    add_fig(doc, "09_steps.png", "图 9  横轴每一档是一次插入。绿=1 次 WQE，橙=2 次 WQE，红=失败")
    add_p(doc, "前七档几乎全绿，橙色只有柱顶一条细边，红色是 0。无副表的失败只有 0.072%，整根柱子看起来仍是绿色，柱顶用红字标出这 7169 条，条数和均匀 1000 万的溢出条数相同。2^22 行的柱顶红字是失败 7.07%。粗哈希的柱顶红字是失败 14.1%，橙色和红色占了绝大部分，因为很多 key 挤进少数行。")
    add_p(doc, "小结：成功的插入只有 1 次或 2 次 WQE。失败只出现在去掉副表、收紧行数、或哈希没有把热点打散的时候。")

    add_p(doc, "图 10 的横轴 k 是哈希落到同一主表行的 key 条数，从 0 画到 12。k=0 是空行，k=4 是刚好写满 4 个槽，k>4 表示这行还有 key 要进冲突表。纵轴是这样的行占主表 2^24 行的比例。实线是仿真，虚线是 Poisson。圆点是 1000 万条，λ = 0.596；方点是配满 1677 万条，λ = 1。")
    add_fig(doc, "10_hist_uniform.png", "图 10  横轴是一行上的到达条数 k。两条实线分别是 1000 万和配满，虚线是对应的 Poisson")
    add_p(doc, "两条实线和虚线重合。配满时 k=0 和 k=1 各约占 37%，k=2 约占 18%，更长的尾巴变矮。这就是主表大约四分之三的槽空着、溢出却只有 0.44% 的原因：大多数行只来了 0 条或 1 条。")
    add_p(doc, "小结：均匀哈希的行到达数服从 Poisson，图上的点和公式是同一条曲线。")

    add_p(doc, "图 11 的横轴四项都按同一批 16,777,216 条均匀 key 计算。从左到右：主表槽位利用率、主表溢出占 key 数的百分比、需要第 2 次 WQE 的 key 占 key 数的百分比、插入失败占 key 数的百分比。每一组左边青色柱是驱动几何 2^24 行，右边珊瑚色柱是紧凑几何 2^22 行。利用率的分母是槽位数 4×行数，后三项的分母是 key 数。")
    add_fig(doc, "12_d_compare.png", "图 11  横轴四项的含义见上。左柱 2^24 行，右柱 2^22 行，key 都是 16,777,216 条")
    add_p(doc, "2^24 行的利用率约 24.9%，溢出和 2 次 WQE 都是 0.44%，失败是 0。2^22 行的利用率升到约 80.5%，溢出约 19.5%，失败约 7.07%。利用率没有到 100%，是因为溢出的 key 没有留在主表槽里。")
    add_p(doc, "小结：多出来的行换来的是失败率从 7% 回到 0，代价是主表大约四分之三的槽空着。")

    add_p(doc, "图 12 的横轴是 λ，等于插入条数除以主表行数。四个位置从左到右：0.596 是 1000 万条除以 2^24 行，1 是配满，1.19 是 2000 万条除以 2^24 行，4 是 1677 万条除以 2^22 行。纵轴是百分比。前三个点连成线，属于驱动几何；λ=4 的两个点单独标出，属于紧凑几何。青色是主表利用率，绿色是驱动几何的失败率，珊瑚色点是紧凑几何的利用率，橙色点是紧凑几何的失败率。")
    add_fig(doc, "13_capacity.png", "图 12  横轴 λ = 条数 / 行数。0.596、1、1.19 是 2^24 行，λ=4 是 2^22 行")
    add_p(doc, "驱动几何上利用率随 λ 近似线性升到约 25%，失败率贴着 0。紧凑几何的利用率跳到约 80%，失败率跳到约 7%。空着的槽是行内容量：每行只放平均 1 条时，Poisson 起伏很少能填满 4 个槽。")
    add_p(doc, "小结：失败率对 λ 很敏感。驱动把配满点放在 λ=1，紧凑几何把同一批 key 放到 λ=4。")

    add_p(doc, "图 13 的横轴 k 仍是落到同一主表行的 key 条数，只画 0 到 20。这张图是粗哈希、1000 万条、主表 2^24 行。纵轴是这样的行占 2^24 行的比例，用对数，否则 k=0 那根约 97.9% 的柱子会把后面的柱子压成一条线。k≥47 的行有 6986 行，没有画在图里；k=47 这个桶在数据里是“47 以及更多”的合计。")
    add_fig(doc, "11_hist_zipf.png", "图 13  横轴 k=0 到 20。纵轴是行的比例，对数。粗哈希、1000 万条。k≥47 的 6986 行不在图上")
    add_p(doc, "k=0 的柱子占了大约 97.9% 的行。少数行的 k 一直延伸到几十。主表利用率因此只有约 1%，热点行立刻溢出，碰撞表也接不住，失败约 14.1%。")
    add_p(doc, "小结：行号如果只由一个重尾字段决定，空行占绝大多数，失败集中在少数热点行。")

    headers = ["实验", "n", "溢出", "公式", "相对误差", "失败", "失败公式", "u_main", "最大比较"]
    body = []
    order = [
        ("A", "uniform", "A 均匀"),
        ("B", "uniform", "B 配满 2^24"),
        ("C", "uniform", "C 2,000 万"),
        ("A", "uniform_alt", "A 另一哈希"),
        ("A", "zipf_mixed", "A 重尾全字段哈希"),
        ("A", "ctu_resample", "A CTU 采样 1,000 万"),
        ("TraceFull", "ctu", "CTU 全量 1,602,033"),
        ("Ablation", "uniform", "消融，无碰撞表"),
        ("Dtight", "uniform", "D-tight 2^22 行"),
        ("A", "zipf_coarse", "A 粗哈希重尾"),
        ("B", "zipf_coarse", "B 粗哈希重尾"),
        ("C", "zipf_coarse", "C 粗哈希重尾"),
    ]
    for exp, data, name in order:
        r = by[(exp, data)]
        body.append([
            name,
            f"{int(r['n']):,}",
            pct(r["over_frac"]),
            pct(r["theory_over"]),
            rel(r["over_frac"], r["theory_over"]),
            pct(r["fail_frac"]),
            pct(r["theory_fail"]),
            pct(r["u_main"]),
            r["lookup_max_cmp"],
        ])
    add_table(doc, headers, body)
    add_p(doc, "表 1  全部实插。u_main 是主表已占用槽除以主表槽位总数。公式失败率为 0 的格子，数值小于 10^-8，打印成 0.0000%。CTU 全量的期望溢出条数约 1 条，仿真为 1 条。", size=9)

    add_p(doc, f"均匀 A 的溢出相对误差 {rel(a['over_frac'], a['theory_over'])}，落在计划要求的 5% 以内。配满点和 2,000 万条同样贴住公式，失败条数都是 0。第二套哈希的溢出是 {pct(by[('A','uniform_alt')]['over_frac'])}，与第一套同一量级。CTU 采样的 1,000 万条溢出 {pct(by[('A','ctu_resample')]['over_frac'])}，相对误差 {rel(by[('A','ctu_resample')]['over_frac'], by[('A','ctu_resample')]['theory_over'])}。真实五元组经过与均匀实验相同的混合之后，行占用回到 Poisson。")
    ab = by[("Ablation", "uniform")]
    add_p(doc, f"消融实验的失败条数 {int(ab['fail']):,}，等于均匀 A 的溢出条数 {int(a['overflow']):,}。碰撞表在驱动几何下承担的是这几千条尾部规则的成功率。步数上界则来自两级深度都有限：即使粗哈希把 1,000 万条里的 {pct(by[('A','zipf_coarse')]['fail_frac'])} 打成失败，成功的那些仍然只走了 1 步或 2 步，查找最大比较次数是 20，也就是 4 + 16。")
    add_p(doc, f"D-tight 的主表利用率 {pct(d['u_main'])}，溢出 {pct(d['over_frac'])}（公式 {pct(d['theory_over'])}），失败 {pct(d['fail_frac'])}（公式 {pct(d['theory_fail'])}）。槽位总数等于 key 数时，哈希方差仍让约 19.5% 的 key 无法留在自己的行里，其中约 7.1% 连深度 16 的碰撞行也放不下。驱动几何在同一批 key 上的主表利用率是 {pct(b['u_main'])}，失败数是 0。多出来的行把 λ 从 4 降到 1。")
    z = by[("A", "zipf_coarse")]
    zm = by[("A", "zipf_mixed")]
    add_p(doc, f"重尾全字段哈希的溢出 {pct(zm['over_frac'])}，与均匀 A 一致，失败为 0。粗哈希下 1,000 万条的主表利用率只有 {pct(z['u_main'])}，因为绝大多数行是空的，热点行立刻溢出；碰撞表利用率 {pct(z['u_col'])}，失败 {int(z['fail']):,} 条。配满和 2,000 万条时碰撞表利用率接近 100%，失败比例升到 {pct(by[('B','zipf_coarse')]['fail_frac'])} 和 {pct(by[('C','zipf_coarse')]['fail_frac'])}。")

    add_h(doc, "7.3  主表负载 25% 到 100%", 2)
    add_p(doc, "负载按插入条数除以 2^24 计算。2^24 是驱动为 num_log = 24 配出的主表行数，100% 就是平均每行 1 条 key。五档的条数是 4,194,304、8,388,608、11,744,051、15,099,494 和 16,777,216。70% 对应 bwc.c 里开始扩行的水位（规则数达到 2^size_log 的 70%）。哈希仍是均匀的，表几何仍是主表每行 4 个 STE、冲突表 2^19 行、每行 16 个 STE。")
    add_p(doc, "主表溢出率是主行 4 个槽放不下、因而改写冲突表的规则占全部插入的比例。冲突表溢出率是冲突行 16 个槽也放不下的规则占全部插入的比例，也就是插入失败率。")
    load_order = [
        ("L25", "25%"), ("L50", "50%"), ("L70", "70%"),
        ("L90", "90%"), ("L100", "100%"),
    ]
    load_rows = []
    for exp, lab in load_order:
        r = by[(exp, "uniform")]
        load_rows.append([
            lab,
            f"{int(r['n']):,}",
            f"{int(r['overflow']):,}",
            pct(r["over_frac"]),
            pct(r["theory_over"]),
            f"{int(r['fail']):,}",
            pct(r["fail_frac"]),
            pct(r["theory_fail"]),
        ])
    add_table(doc,
              ["负载", "插入条数", "主表溢出条数", "主表溢出率", "公式",
               "冲突表溢出条数", "冲突表溢出率（失败率）", "失败公式"],
              load_rows)
    add_p(doc, "表 2  均匀哈希、驱动几何。冲突表溢出率等于插入失败率。", size=9)
    add_p(doc, "图 14 的横轴五档都是插入条数 / 2^24，主表始终是 2^24 行 × 4。从左到右：25% = 4,194,304 条，50% = 8,388,608 条，70% = 11,744,051 条，90% = 15,099,494 条，100% = 16,777,216 条。70% 是 BWC 扩行水位。左图纵轴是主表溢出占该档条数的百分比，右图纵轴是冲突表溢出，也就是插入失败率。")
    add_fig(doc, "14_load_overflow.png", "图 14  横轴是条数 / 2^24 的五档。左为主表溢出率，右为冲突表溢出率")
    add_p(doc, "左图从 25% 的 124 条（0.0030%）升到 100% 的 73,470 条（0.4379%），仿真柱和公式柱重叠。25% 这一档期望溢出只有一百条出头，相对误差会显得大一些；50% 以上相对误差回到 2% 以内。右图五档都是 0。100% 并不是槽位填满：这一档平均每行只有 1 条 key，槽位利用率约 25%。")
    add_p(doc, "小结：驱动几何从 25% 加到配满，主表溢出增加，冲突表一次都没有溢出。")
    add_p(doc, "图 15 的横轴是主表一行里已经写下的 STE 个数，只有 0、1、2、3、4。0 是空行，4 是四个槽都满，更长的到达被记到冲突表，所以横轴到 4 为止。同一个数字上从左到右五根柱是 25%、50%、70%、90%、100%，条数和图 14 相同。纵轴是这样的行占主表 2^24 行的百分比。")
    add_fig(doc, "15_main_row_depth.png", "图 15  横轴是主表一行已占用的 STE 数，0 到 4。五根柱从左到右是 25% 到 100%")
    add_p(doc, "25% 时深度 0 和 1 最高。负载升高后深度 0 变矮，深度 2、3、4 变高。100% 时深度 0 占 36.78%，1 占 36.80%，2 占 18.39%，3 占 6.13%，4 占 1.90%，与 λ=1 的 Poisson 在 0 到 3 的概率以及 P(X≥4) 对齐。")
    add_p(doc, "小结：配满时只有约 1.9% 的行写满 4 个槽，所以大多数 key 留在主表。")
    add_p(doc, "图 16 的横轴是冲突表一行里已经写下的 STE 个数，从 0 到 16。五条线仍是图 14 的五档负载。纵轴是这样的行占冲突表 2^19 行的百分比，用对数，因为空行远多于有 key 的行。")
    add_fig(doc, "16_col_row_depth.png", "图 16  横轴是冲突表一行已占用的 STE 数，0 到 16。纵轴为行的百分比，对数坐标")
    add_p(doc, "五条线都在深度 4 之后掉到图外，深度 5 到 16 的行数是 0。100% 时 86.92% 的冲突行深度为 0，12.19% 为 1，0.85% 为 2，最深只到 4。16 个槽的上限没有被碰到。")
    add_p(doc, "小结：驱动几何的冲突表在配满时仍然很空，所以插入失败率是 0。")

    add_h(doc, "7.4  2^22 行的对照", 2)
    add_p(doc, "对照表把主表收成 2^22 行、每行仍是 4 个槽，冲突表按同样的规则减 5，得到 2^17 行、每行 16 个槽。五档使用和上一节相同的 key 条数，负载仍按 n / 2^24 标注。100% 就是把 2^24 条 key 放进 2^22 行，平均每行 4 条，槽位总数等于 key 数。这是实验 D-tight 的同一张表，这里把它补到 25%、50%、70%、90%、100% 五档，并给出每行深度。")
    tight_rows = []
    for (wexp, lab), (texp, _tlab) in zip(load_order, [
        ("D25", "25%"), ("D50", "50%"), ("D70", "70%"),
        ("D90", "90%"), ("D100", "100%"),
    ]):
        w = by[(wexp, "uniform")]
        t = by[(texp, "uniform")]
        tight_rows.append([
            lab,
            f"{int(w['n']):,}",
            pct(w["over_frac"]),
            pct(t["over_frac"]),
            pct(t["theory_over"]),
            pct(w["fail_frac"]),
            pct(t["fail_frac"]),
            pct(t["theory_fail"]),
        ])
    add_table(doc,
              ["负载", "插入条数", "2^24 主表溢出", "2^22 主表溢出", "2^22 溢出公式",
               "2^24 失败率", "2^22 失败率", "2^22 失败公式"],
              tight_rows)
    add_p(doc, "表 3  同一批均匀 key。2^22 行的失败率就是这张紧凑表的冲突表溢出率。", size=9)
    add_p(doc, "图 17 的横轴和图 14 相同，仍是条数 / 2^24 的五档，条数也相同。变的是表：主表 2^22 行 × 4，冲突表 2^17 行 × 16。所以 100% 这一档平均每行是 4 条，不是 1 条。左图纵轴是主表溢出率，右图纵轴是冲突表溢出率，也就是插入失败率。")
    add_fig(doc, "17_tight_overflow.png", "图 17  横轴仍是条数 / 2^24。表已收成 2^22 行。左为溢出，右为失败")
    add_p(doc, "25% 时这张小表的 λ 已经是 1，溢出 0.4310%，和驱动几何配满时的 0.4379% 是同一量级，失败是 0。50% 溢出 3.75%，失败仍是 0。70% 第一次出现失败，989 条，0.0084%。90% 失败 2.63%，100% 溢出 19.53%、失败 7.07%。每一档都和 Poisson 公式重合。")
    add_p(doc, "小结：同一批 key 放到 2^22 行上，50% 以前冲突表还接得住，70% 开始失败，配满时失败约 7%。")
    add_p(doc, "图 18 的横轴仍是那五档条数 / 2^24。左图纵轴是主表溢出率，用对数，因为 2^24 行的溢出是千分之几，2^22 行是百分之几到百分之二十。右图纵轴是插入失败率，线性坐标。青色是 2^24 行，珊瑚色是 2^22 行。")
    add_fig(doc, "18_geometry_compare.png", "图 18  横轴是同一批 key 的五档。青线 2^24 行，红线 2^22 行。左为溢出（对数），右为失败率")
    add_p(doc, "两条溢出曲线一起上升，2^22 行始终高一个多数量级。失败率只有红线在 70% 之后离开 0，青线五档都贴在 0 上。")
    add_p(doc, "小结：行数从 2^24 收到 2^22，溢出和失败都变大；驱动几何的失败率在这五档保持为 0。")
    add_p(doc, "图 19 的横轴是 2^22 行主表里一行写下的 STE 个数，0 到 4。同一个数字上从左到右五根柱仍是 25% 到 100%。纵轴是这样的行占这 2^22 行的百分比。深度 4 表示四个槽都满，多出来的 key 在冲突表。")
    add_fig(doc, "19_tight_main_depth.png", "图 19  横轴是 2^22 行主表的行深，0 到 4。五根柱从左到右是 25% 到 100%")
    add_p(doc, "25% 的形状和驱动几何配满时一样，因为两边的 λ 都是 1。100% 时深度 4 的行占 56.66%，空行只剩 1.84%。一半以上的行已经写满，所以约五分之一的 key 进了冲突表。")
    add_p(doc, "小结：2^22 行在 100% 负载时，主表多数行是满的。")
    add_p(doc, "图 20 的横轴是这张紧凑表的冲突行深度，0 到 16。冲突表有 2^17 行。五条线仍是 25% 到 100%。纵轴是这样的行占 2^17 行的百分比，对数坐标。深度 16 表示 16 个槽都满，再来的 key 就是插入失败。")
    add_fig(doc, "20_tight_col_depth.png", "图 20  横轴是 2^17 行冲突表的行深，0 到 16。深度 16 表示该行已满")
    add_p(doc, "25% 最深到 4，50% 最深到 12，这两档没有深度 16 的行，失败是 0。70% 第一次在深度 16 出现一截，占冲突行的 0.94%，对应 989 条失败。90% 有 73.40% 的冲突行停在 16。100% 有 97.83% 的冲突行停在 16，失败率 7.07%。")
    add_p(doc, "小结：紧凑几何的失败从冲突行顶满 16 个槽开始。驱动几何的冲突行最深只到 4，所以没有这一截。")
    add_p(doc, "十三张实验图合在一起：驱动几何把配满点放在平均每行 1 条。均匀哈希、换一套混合、以及 CTU 五元组，溢出都停在公式给出的千分之几，冲突表接得住，失败是 0，成功插入只有 1 次或 2 次 WQE。同一批 key 收成 2^22 行之后，失败率升到约 7%，深度图能看到冲突行顶满。粗哈希则让少数行先满，空行占绝大多数，失败率离开设计点。")

    add_h(doc, "8  机制在什么范围内站得住", 1)
    add_p(doc, "均匀哈希、以及能把五元组打散的哈希，在 num_log = 24 的表上把 2,000 万条放完仍然失败 0 条。第二步 WQE 的比例从 1,000 万条时的 0.07% 升到配满时的 0.44%，再升到 2,000 万条时的 0.78%。插入代价保持为一次或两次定长设备操作。这就是“表已经有千万条，再插一条仍然是常数时间”的硬件原因：哈希只选一行，行深是常数，满了还有一张行深同样有常数上界的副表。")
    add_p(doc, "主表 75% 左右的空槽是这个常数上界的代价。BWC 还在 70% 行水位就扩行，所以线上常用水位比 λ = 1 更松，STE 相对规则数的余量比配满点更大。")
    add_p(doc, "上界先失效的方式有两种。一是碰撞行也被占满，插入从“最多两次 WQE”变成失败，粗哈希重尾和 D-tight 都走到了这里。二是 key 本身很短，许多互异规则落在同一行，副表接不住热点。步数上界对成功的插入仍然成立，先丢的是成功率。")
    add_p(doc, "模型没有覆盖的部分：设备内部哈希多项式的具体系数、GTA 引擎每秒能退休多少 WQE、以及 STC 动作链本身的延迟。那些决定 R_wqe 的具体数值。本文证明的是驱动选的行数和深度，在均匀与真实五元组哈希下，把溢出和失败控制在公式给出的位置，并且把单次插入限制在两条 WQE 以内。")

    add_h(doc, "9  复现", 1)
    add_p(doc, "在 mlx5-hws-analysis/sim 下执行 make，然后 ./hws_sim out。均匀、重尾、D-tight 和消融不需要外部文件。CTU 原始文件与采样文件通过第 3、第 4 个参数传入。画图使用 draw_figures.py，本文由 build_docx.py 生成。")
    add_p(doc, "仿真输出在 sim/out/summary.csv、hist.csv 和 depth.csv。depth.csv 是主表和冲突表每一档占用深度的行数。各次运行的墙钟时间在 0.5 秒到 8 秒之间，包含插入和全量查找，只说明这个用户态模型在单线程下的速度。")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
